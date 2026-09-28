from __future__ import annotations

import json
import os
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
INPUT_FILE = ROOT / "input" / "schedule.xlsx"
CONFIG_FILE = ROOT / "input" / "config.json"
OUTPUT_FILE = ROOT / "schedule.ics"


def is_day_number(value: object) -> bool:
    try:
        number = int(value)
    except (TypeError, ValueError):
        return False
    return number == float(value) and 1 <= number <= 31


def day_from_value(value: object) -> int | None:
    if isinstance(value, datetime):
        return value.day
    if isinstance(value, date):
        return value.day
    if is_day_number(value):
        return int(value)

    text = str(value or "").strip()
    match = re.search(r"(?:^|[-/])(\d{1,2})$", text)
    if match:
        day = int(match.group(1))
        return day if 1 <= day <= 31 else None
    return None


def find_header(rows: list[list[object]]) -> tuple[int, list[int]]:
    for row_index in range(min(8, len(rows))):
        columns = [
            index
            for index, value in enumerate(rows[row_index])
            if day_from_value(value) is not None
        ]
        if len(columns) >= 7:
            return row_index, columns
    raise ValueError("没有找到日期表头")


def find_name_column(rows: list[list[object]], header_row: int) -> int:
    for row_index in range(header_row + 1):
        for column_index, value in enumerate(rows[row_index]):
            if "姓名" in str(value or "") or "名字" in str(value or ""):
                return column_index
    return 3


def escape_ics(value: object) -> str:
    return (
        str(value)
        .replace("\\", "\\\\")
        .replace(";", "\\;")
        .replace(",", "\\,")
        .replace("\r", "")
        .replace("\n", "\\n")
    )


def fold_line(line: str) -> list[str]:
    result: list[str] = []
    current = ""
    current_bytes = 0
    for character in line:
        character_bytes = len(character.encode("utf-8"))
        if current and current_bytes + character_bytes > 74:
            result.append(current)
            current = f" {character}"
            current_bytes = 1 + character_bytes
        else:
            current += character
            current_bytes += character_bytes
    if current:
        result.append(current)
    return result


def ics_date(value: date) -> str:
    return value.strftime("%Y%m%d")


def main() -> None:
    if not INPUT_FILE.exists():
        raise FileNotFoundError("找不到 input/schedule.xlsx")
    if not CONFIG_FILE.exists():
        raise FileNotFoundError("找不到 input/config.json")

    config = json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
    target_name = str(config.get("name", "")).strip()
    if not target_name:
        raise ValueError("input/config.json 缺少 name")

    workbook = load_workbook(INPUT_FILE, data_only=True)
    sheet = workbook.active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    header_row, date_columns = find_header(rows)
    name_column = find_name_column(rows, header_row)

    today = datetime.now(ZoneInfo("Asia/Shanghai")).date()
    first_day = day_from_value(rows[header_row][date_columns[0]])
    if first_day is None:
        raise ValueError("无法读取第一个日期")
    first_date = date(today.year, today.month, first_day)

    target_row = None
    for row_index in range(header_row + 1, len(rows)):
        name = str(rows[row_index][name_column] or "").strip()
        if name == target_name:
            target_row = row_index
            break
    if target_row is None:
        raise ValueError(f"找不到姓名：{target_name}")

    stamp = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y%m%dT%H%M%S")
    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//Shift Calendar Exporter//CN",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
        "X-WR-CALNAME:排班日历",
        "X-WR-TIMEZONE:Asia/Shanghai",
        "X-PUBLISHED-TTL:PT1D",
        "REFRESH-INTERVAL;VALUE=DURATION:P1D",
    ]

    for offset, column_index in enumerate(date_columns):
        current_date = first_date + timedelta(days=offset)
        code = str(rows[target_row][column_index] or "").strip()
        if not code:
            continue
        hours = rows[target_row + 1][column_index] if target_row + 1 < len(rows) else None
        description = ["来源：排班表"]
        if isinstance(hours, (int, float)) and hours > 0:
            description.append(f"工时：{hours}小时")
        next_date = current_date + timedelta(days=1)

        lines.extend(
            [
                "BEGIN:VEVENT",
                f"UID:shift-{ics_date(current_date)}-{offset}@calendar.local",
                f"DTSTAMP:{stamp}",
                f"DTSTART;VALUE=DATE:{ics_date(current_date)}",
                f"DTEND;VALUE=DATE:{ics_date(next_date)}",
                f"SUMMARY:{escape_ics(code)}",
                f"DESCRIPTION:{escape_ics(chr(10).join(description))}",
                "TRANSP:TRANSPARENT",
                "END:VEVENT",
            ]
        )

    lines.append("END:VCALENDAR")
    OUTPUT_FILE.write_text(
        "\r\n".join(line for item in lines for line in fold_line(item)) + "\r\n",
        encoding="utf-8",
    )
    print(f"Generated {OUTPUT_FILE} for {target_name}")


if __name__ == "__main__":
    main()
